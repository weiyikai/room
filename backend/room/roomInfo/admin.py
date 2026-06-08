from import_export import resources
from roomInfo.models import DictType, DictItem, RoomInfo
from django.contrib import admin
from django.http import HttpResponse
from django.conf import settings
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
import os
import traceback

# 顶部必须添加这两个导入（放到文件最上面）
from io import BytesIO
from openpyxl.drawing.image import Image

import datetime  # 新增：生成时间戳

# 注册字典分类
@admin.register(DictType)
class DictTypeAdmin(admin.ModelAdmin):
    # 👇 只在最前面添加 id 字段，其余保持你原来的不变
    list_display = ('id', 'name', 'code', 'remark', 'created_at')
    search_fields = ('name', 'code')
    list_filter = ('created_at',)
    ordering = ('code',)
    # 👇 新增：设置ID为只读，不可编辑
    readonly_fields = ('id',)

# 注册字典项
@admin.register(DictItem)
class DictItemAdmin(admin.ModelAdmin):
    list_display = ('id','label', 'value', 'type', 'parent', 'sort', 'status')
    search_fields = ('label', 'value')
    list_filter = ('type', 'status')
    ordering = ('type', 'sort')
    # 👇 新增：设置ID为只读，不可编辑
    readonly_fields = ('id',)

# 导出资源配置
class RoomInfoResource(resources.ModelResource):
    class Meta:
        model = RoomInfo
        fields = '__all__'  # 导出所有字段


# -------------------------- 【终极根治版：强制读数据库最新数据，无缓存】 --------------------------
def export_room_info_with_images(modeladmin, request, queryset):
    # 1. 【核心】获取前台筛选/排序/勾选后的 主键ID列表（保留所有前台条件）
    selected_ids = queryset.values_list('id', flat=True)
    print(f"📊 前台筛选后的主键ID：{list(selected_ids)}")

    # 2. 【你要的逻辑】根据 主键ID 去数据库实时查询最新数据！
    # 严格按RoomInfo主键ID查询，绕过缓存，只查选中/筛选的条目
    queryset = RoomInfo.objects.filter(id__in=selected_ids).order_by('id')

    # 调试：打印最新总条数
    print(f"📊 数据库最新总条数：{queryset.count()}")

    # 1. 创建工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "机房信息(含图片)"

    # 2. 表头
    headers = [
        "ID", "所属省份", "所属地市", "所属区县", "局站名称", "局站等级",
        "机房名称", "机房图片", "机房长度", "机房长度图片", "机房宽度", "机房宽度图片", "机房高度", "机房高度图片",
        "网络机柜数", "网络机柜图片", "ODF/MODF数量", "ODF/MODF图片",
        "电源列头柜数量", "电源列头柜图片", "开关电源数量", "开关电源图片", "开关电源型号", "开关电源型号图片",
        "开关电源容量", "开关电源容量图片", "开关电源厂家", "开关电源厂家图片", "开关电源出厂日期", "出厂日期图片",
        "开关电源模块型号", "模块型号图片", "开关电源模块容量", "模块容量图片", "开关电源模块数量", "模块数量图片",
        "电池型号", "电池型号图片", "电池容量", "电池容量图片", "电池厂家", "电池厂家图片", "电池组数", "电池组图片",
        "空调型号", "空调型号图片", "空调容量", "空调容量图片", "空调厂家", "空调厂家图片", "空调数量", "空调数量图片",
        "油机型号", "油机型号图片", "油机容量", "油机容量图片", "油机厂家", "油机厂家图片", "油机数量", "油机数量图片",
        "交流配电箱型号", "配电箱型号图片", "交流配电箱容量", "配电箱容量图片", "交流配电箱厂家", "配电箱厂家图片",
        "交流配电箱数量", "配电箱数量图片"
    ]
    ws.append(headers)

    # 3. 图片配置
    IMG_W = 120
    IMG_H = 80

    # 4. 遍历导出（稳定无丢失）
    for row_num, room in enumerate(queryset, start=2):
        print(f"➡️ 导出机房ID：{room.id}")
        ws.row_dimensions[row_num].height = IMG_H + 10

        # 写入数据
        row_data = [
            # 基础信息
            room.id,
            room.province.label if room.province else "",
            room.city.label if room.city else "",
            room.area.label if room.area else "",
            room.station.label if room.station else "",
            room.station_level.label if room.station_level else "",

            # 机房基础参数
            room.room_name or "","",
            room.room_length or "","",
            room.room_width or "","",
            room.room_height or "","",

            # 网络设备
            room.cabinet_count or "","",
            room.odf_count or "","",

            # 开关电源
            room.power_cabinet_count or "","",
            room.switch_power_count or "","",
            room.switch_power_model or "","",
            room.switch_power_capacity or "","",
            room.switch_power_factory or "","",
            room.switch_power_date or "","",

            # 电池设备
            room.switch_module_model or "","",
            room.switch_module_capacity or "","",
            room.switch_module_count or "","",
            room.battery_model or "","",
            room.battery_capacity or "","",
            room.battery_factory or "","",
            room.battery_group_count or "","",

            # 空调
            room.ac_model or "","",
            room.ac_capacity or "","",
            room.ac_factory or "","",
            room.ac_count or "","",

            # 油机
            room.generator_model or "","",
            room.generator_capacity or "","",
            room.generator_factory or "","",
            room.generator_count or "","",

            # 交流配电箱
            room.ac_box_model or "","",
            room.ac_box_capacity or "","",
            room.ac_box_factory or "","",
            room.ac_box_count or "","",

            # 时间
        ]
        ws.append(row_data)

        # 机房基础参数图片
        _insert_img(ws, room.room_name_img, row_num, 8, IMG_W, IMG_H)
        _insert_img(ws, room.room_length_img, row_num, 10, IMG_W, IMG_H)
        _insert_img(ws, room.room_width_img, row_num, 12, IMG_W, IMG_H)
        _insert_img(ws, room.room_height_img, row_num, 14, IMG_W, IMG_H)

        # 网络设备图片
        _insert_img(ws, room.cabinet_count_img, row_num, 16, IMG_W, IMG_H)
        _insert_img(ws, room.odf_count_img, row_num, 18, IMG_W, IMG_H)

        # 开关电源图片
        _insert_img(ws, room.power_cabinet_count_img, row_num, 20, IMG_W, IMG_H)
        _insert_img(ws, room.switch_power_count_img, row_num, 22, IMG_W, IMG_H)
        _insert_img(ws, room.switch_power_model_img, row_num, 24, IMG_W, IMG_H)
        _insert_img(ws, room.switch_power_capacity_img, row_num, 26, IMG_W, IMG_H)
        _insert_img(ws, room.switch_power_factory_img, row_num, 28, IMG_W, IMG_H)
        _insert_img(ws, room.switch_power_date_img, row_num, 30, IMG_W, IMG_H)

        # 电池设备图片
        _insert_img(ws, room.switch_module_model_img, row_num, 32, IMG_W, IMG_H)
        _insert_img(ws, room.switch_module_capacity_img, row_num, 34, IMG_W, IMG_H)
        _insert_img(ws, room.switch_module_count_img, row_num, 36, IMG_W, IMG_H)
        _insert_img(ws, room.battery_model_img, row_num, 38, IMG_W, IMG_H)
        _insert_img(ws, room.battery_capacity_img, row_num, 40, IMG_W, IMG_H)
        _insert_img(ws, room.battery_factory_img, row_num, 42, IMG_W, IMG_H)
        _insert_img(ws, room.battery_group_count_img, row_num, 44, IMG_W, IMG_H)

        # 空调图片
        _insert_img(ws, room.ac_model_img, row_num, 46, IMG_W, IMG_H)
        _insert_img(ws, room.ac_capacity_img, row_num, 48, IMG_W, IMG_H)
        _insert_img(ws, room.ac_factory_img, row_num, 50, IMG_W, IMG_H)
        _insert_img(ws, room.ac_count_img, row_num, 52, IMG_W, IMG_H)

        # 油机图片
        _insert_img(ws, room.generator_model_img, row_num, 54, IMG_W, IMG_H)
        _insert_img(ws, room.generator_capacity_img, row_num, 56, IMG_W, IMG_H)
        _insert_img(ws, room.generator_factory_img, row_num, 58, IMG_W, IMG_H)
        _insert_img(ws, room.generator_count_img, row_num, 60, IMG_W, IMG_H)

        # 交流配电箱图片
        _insert_img(ws, room.ac_box_model_img, row_num, 62, IMG_W, IMG_H)
        _insert_img(ws, room.ac_box_capacity_img, row_num, 64, IMG_W, IMG_H)
        _insert_img(ws, room.ac_box_factory_img, row_num, 66, IMG_W, IMG_H)
        _insert_img(ws, room.ac_box_count_img, row_num, 68, IMG_W, IMG_H)

    # 列宽
    for col in range(1, 69):
        ws.column_dimensions[get_column_letter(col)].width = 20

    # ====================== 修复1：改用BytesIO保存，解决文件流报错 ======================
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    # 生成当前时间戳（每次导出都不一样）
    timestamp = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
    # 文件名：机房信息_带图片_20260424162030.xlsx
    filename = f"机房信息_带图片_{timestamp}.xlsx"

    # 双重禁缓存 + 唯一文件名
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    # 禁止浏览器缓存
    response['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    response['Pragma'] = 'no-cache'
    response['Expires'] = '0'
    # 🔥 核心：唯一文件名，彻底杜绝缓存
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    return response


# -------------------------- 【✅ 永久稳定版：解决偶发无图 + 文件流关闭报错】 --------------------------
def _insert_img(ws, img_field, row, col, width, height):
    try:
        if not img_field:
            return

        # 你的配置是对的！直接用 MEDIA_ROOT 拼接，不动
        img_path = os.path.abspath(os.path.join(settings.MEDIA_ROOT, img_field.name))

        # 检查文件是否存在
        if not os.path.exists(img_path):
            print(f"[无图] 文件不存在：{img_path}")
            return

        # 🔥 核心修复：将图片数据完整读入内存BytesIO，彻底解决文件流关闭报错
        with open(img_path, "rb") as f:
            img_bytes = BytesIO(f.read())

        img = Image(img_bytes)
        # 设置图片大小
        img.width = width
        img.height = height
        # 插入单元格
        cell_pos = f"{get_column_letter(col)}{row}"
        ws.add_image(img, cell_pos)

        print(f"[成功插入] 图片：{img_field.name}")

    except Exception as e:
        print(f"[插入失败] 原因：{str(e)}")
        traceback.print_exc()


export_room_info_with_images.short_description = "✅ 导出机房信息（带图片）"

# -------------------------- 【你所有原有配置 100% 保留，仅恢复只读字段】 --------------------------
@admin.register(RoomInfo)
class RoomInfoAdmin(admin.ModelAdmin):
    actions = [export_room_info_with_images]
    list_display = [field.name for field in RoomInfo._meta.fields]
    search_fields = ['station__label', 'area__label', 'city__label']
    list_filter = ['province', 'city', 'area', 'station_level']
    # 恢复图片字段只读（更规范，不影响功能）
    # readonly_fields = [f.name for f in RoomInfo._meta.fields if 'img' in f.name] + ['created_at', 'updated_at']

    # 🔥 新增这一行：固定默认排序，新增数据立刻能导出
    # 202604170922 解决新增数据后全部导出没有这条数据的问题，得排个序刷新一下数据顺序，新增的数据才会被纳入导出的查询集
    # ordering = ['-id']

    # 你原来的 fieldsets 完全不变！
    fieldsets = (
        ("基础信息", {'fields': ('province', 'city', 'area', 'station_level','station')}),
        ("机房基础参数", {'fields': (('room_name', 'room_name_img'), ('room_length', 'room_length_img'),
                                     ('room_width', 'room_width_img'), ('room_height', 'room_height_img'))}),
        ("网络设备", {'fields': (('cabinet_count', 'cabinet_count_img'), ('odf_count', 'odf_count_img'))}),
        ("开关电源",
         {'fields': (('power_cabinet_count','power_cabinet_count_img'),('switch_power_count', 'switch_power_count_img'), ('switch_power_model', 'switch_power_model_img'),
                     ('switch_power_capacity', 'switch_power_capacity_img'),
                     ('switch_power_factory', 'switch_power_factory_img'),
                     ('switch_power_date','switch_power_date_img'))}),
        ("电池设备", {'fields': (('switch_module_model','switch_module_model_img'),
                                 ('switch_module_capacity','switch_module_capacity_img'),
                                 ('switch_module_count','switch_module_count_img'),
                                 ('battery_model', 'battery_model_img'),
                                 ('battery_capacity', 'battery_capacity_img'),
                                 ('battery_factory', 'battery_factory_img'),
                                 ('battery_group_count', 'battery_group_count_img'))}),
        ("空调", {'fields': (('ac_model', 'ac_model_img'),
                             ('ac_capacity','ac_capacity_img'),
                             ('ac_factory','ac_factory_img'),
                             ('ac_count','ac_count_img'))}),
        ("油机", {'fields': (('generator_model', 'generator_model_img'),
                             ('generator_capacity','generator_capacity_img'),
                             ('generator_factory','generator_factory_img'),
                             ('generator_count','generator_count_img'))}),
        ("交流配电箱", {'fields': (('ac_box_model', 'ac_box_model_img'),
                                    ('ac_box_capacity', 'ac_box_capacity_img'),
                                    ('ac_box_factory', 'ac_box_factory_img'),
                                    ('ac_box_count', 'ac_box_count_img'))}),
    )


# ====================== 巡检轨迹管理 Admin ======================
from .models import InspectionSession, InspectionTrackPoint


def export_inspection_track_excel(modeladmin, request, queryset):
    """导出选中的巡检轨迹为 Excel（含完整轨迹点）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "巡检轨迹汇总"

    # 表头
    headers = [
        "会话ID", "管理员", "手机号", "巡检局站", "开始时间",
        "结束时间", "状态", "轨迹点数", "备注"
    ]
    ws.append(headers)

    for session in queryset.order_by('-start_time'):
        ws.append([
            session.id,
            session.admin_user.name or "",
            session.admin_user.phone,
            session.station.label if session.station else "",
            session.start_time.strftime("%Y-%m-%d %H:%M:%S") if session.start_time else "",
            session.end_time.strftime("%Y-%m-%d %H:%M:%S") if session.end_time else "",
            session.get_status_display(),
            session.track_points.count(),
            session.notes or "",
        ])

    # 第二个sheet：轨迹点明细
    ws2 = wb.create_sheet("轨迹点明细")
    ws2.append(["会话ID", "管理员", "局站", "序号", "纬度", "经度", "精度(m)", "海拔(m)", "定位时间"])

    for session in queryset.order_by('-start_time'):
        for idx, pt in enumerate(session.track_points.all(), 1):
            ws2.append([
                session.id,
                session.admin_user.name or session.admin_user.phone,
                session.station.label if session.station else "",
                idx,
                pt.latitude,
                pt.longitude,
                pt.accuracy or "",
                pt.altitude or "",
                pt.timestamp.strftime("%Y-%m-%d %H:%M:%S") if pt.timestamp else "",
            ])

    # 列宽
    for col in range(1, 10):
        ws.column_dimensions[get_column_letter(col)].width = 22
    for col in range(1, 10):
        ws2.column_dimensions[get_column_letter(col)].width = 20

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
    filename = f"巡检轨迹报表_{timestamp}.xlsx"

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    response['Pragma'] = 'no-cache'
    response['Expires'] = '0'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


export_inspection_track_excel.short_description = "📊 导出巡检轨迹报表（含轨迹点明细）"


@admin.register(InspectionSession)
class InspectionSessionAdmin(admin.ModelAdmin):
    list_display = [
        'id', 'admin_user', 'station', 'start_time', 'end_time',
        'status', 'point_count', 'created_at'
    ]
    list_filter = ['status', 'station', 'start_time']
    search_fields = ['admin_user__phone', 'admin_user__name', 'station__label']
    ordering = ['-start_time']
    actions = [export_inspection_track_excel]
    readonly_fields = ['id', 'created_at']
    date_hierarchy = 'start_time'

    fieldsets = (
        ("基础信息", {'fields': ('admin_user', 'station')}),
        ("巡检时间", {'fields': ('start_time', 'end_time', 'status')}),
        ("备注", {'fields': ('notes',)}),
    )


@admin.register(InspectionTrackPoint)
class InspectionTrackPointAdmin(admin.ModelAdmin):
    list_display = [
        'id', 'session', 'latitude', 'longitude',
        'accuracy', 'altitude', 'timestamp'
    ]
    list_filter = ['session', 'timestamp']
    search_fields = ['session__admin_user__phone', 'session__station__label']
    ordering = ['session', 'timestamp']
    readonly_fields = ['id', 'created_at']